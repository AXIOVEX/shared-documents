"""Build beginning-ai-course.pptx from the rendered slide PNGs in validate/.

Reads slides/deck.json for slide order and speaker_notes.json for notes.
Run from this directory:  python build_pptx_with_notes.py
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu
from PIL import Image
import json
import re

project = Path(__file__).resolve().parent
validate = project / 'validate'
manifest = json.loads((project / 'slides' / 'deck.json').read_text())
notes = json.loads((project / 'speaker_notes.json').read_text())

pngs = sorted(validate.glob('page-*.png'),
              key=lambda p: int(re.search(r'(\d+)$', p.stem).group(1)))
assert len(pngs) == len(manifest['slides']), \
    f"slide/image count mismatch: {len(pngs)} pngs, {len(manifest['slides'])} manifest entries"

with Image.open(pngs[0]) as im:
    w, h = im.size

prs = Presentation()
prs.core_properties.title = 'AI 101: Beyond the Chatbot'
prs.core_properties.subject = 'AI 101: Beyond the Chatbot - course slide deck'
prs.core_properties.author = 'Axiovex Systems, LLC'
prs.slide_width = Emu(12192000)
prs.slide_height = Emu(round(12192000 * h / w))
blank = next(x for x in prs.slide_layouts if x.name == 'Blank')

for png, item in zip(pngs, manifest['slides']):
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(str(png), 0, 0,
                             width=prs.slide_width, height=prs.slide_height)
    sid = item['id']
    slide.notes_slide.notes_text_frame.text = notes[sid]

out = project / 'beginning-ai-course.pptx'
prs.save(out)
print(out, len(prs.slides), 'slides')
